#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""内容完整性校验（按文本框逐条比对，避免跨框拼接造成的误报）。"""
import glob
import html as ihtml
import io
import os
import re
import sys

from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAGENUM = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
# 纯装饰/符号类文本，不参与比对
NOISE = re.compile(r"^[\s\W_]*$|^[\s]*[↑↓→←▲▼◀▶●○※·—\-–_=+*#]+[\s]*$")

srcs = [f for f in sorted(glob.glob(os.path.join(ROOT, "**", "*.pptx"), recursive=True))
        if "_build" not in f]

total_bad = 0
for src in srcs:
    out = os.path.splitext(src)[0] + ".html"
    prs = Presentation(src)

    units = []          # 每个"应出现"的文本单元
    for s in prs.slides:
        for sh in s.shapes:
            if getattr(sh, "has_table", False):
                for tr in sh.table.rows:
                    for c in tr.cells:
                        t = c.text.strip()
                        if len(t) >= 2 and not PAGENUM.match(t) and not NOISE.match(t):
                            units.append(t)
            elif sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if len(t) >= 2 and not PAGENUM.match(t) and not NOISE.match(t):
                        units.append(t)

    h = open(out, encoding="utf-8").read()
    body = re.sub(r"<style.*?</style>|<script.*?</script>", "", h, flags=re.S)
    body = ihtml.unescape(re.sub(r"<[^>]+>", "", body))
    flat = re.sub(r"\s+", "", body)

    missing = [u for u in units if re.sub(r"\s+", "", u) not in flat]

    flag = "OK " if not missing else "WARN"
    if missing:
        total_bad += 1
    print(f"[{flag}] {os.path.basename(out)[:36]:38s} 文本单元={len(units):5d} "
          f"缺失={len(missing):3d} ({len(missing)/max(len(units),1):.1%})")
    for m in missing[:5]:
        print(f"        缺: {m[:70]}")

print(f"\n存在内容缺失的文件: {total_bad} / {len(srcs)}")