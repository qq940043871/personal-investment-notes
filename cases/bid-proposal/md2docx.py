# -*- coding: utf-8 -*-
"""将技术方案 markdown 转为带封面/目录/样式的 docx。"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"d:/ai_show/hello_my_boy/投标技术方案/01-技术方案正文.md"
OUT = r"d:/ai_show/hello_my_boy/投标技术方案/大宗商品交易平台投标技术方案.docx"

doc = Document()

# 默认字体（中英文）
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_heading_font(p, size, bold=True, color=None):
    for run in p.runs:
        run.font.name = '黑体'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        if color:
            run.font.color.rgb = color

def add_field(paragraph, instr):
    r = paragraph.add_run()
    fldBegin = OxmlElement('w:fldChar'); fldBegin.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = instr
    fldSep = OxmlElement('w:fldChar'); fldSep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "右键更新目录"
    fldEnd = OxmlElement('w:fldChar'); fldEnd.set(qn('w:fldCharType'), 'end')
    r._r.append(fldBegin); r._r.append(instrText); r._r.append(fldSep); r._r.append(t); r._r.append(fldEnd)

def add_runs(p, text):
    """处理 **加粗** 行内格式。"""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2]); run.bold = True
        elif part:
            p.add_run(part)

# 封面
cover_lines = [
    ("大宗商品交易供应链平台建设项目", 22, True),
    ("投标技术方案", 16, True),
    ("", 10, False),
    ("投标单位：XX科技有限公司", 12, False),
    ("项目名称：XX（招标方）大宗商品交易供应链平台建设", 12, False),
    ("文档版本：V1.0", 12, False),
    ("编制日期：2026 年 XX 月 XX 日", 12, False),
    ("密级：投标机密 · 未经授权不得扩散", 12, False),
]
for txt, sz, bold in cover_lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(txt)
    run.font.size = Pt(sz); run.bold = bold
    run.font.name = '黑体'; run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

doc.add_page_break()

# 目录
toc_p = doc.add_paragraph(); toc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = toc_p.add_run("目  录"); tr.bold = True; tr.font.size = Pt(16)
tr.font.name = '黑体'; tr.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
doc.add_paragraph()
add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
doc.add_page_break()

# 解析正文
with open(SRC, encoding='utf-8') as f:
    lines = f.read().split('\n')

i = 0
n = len(lines)
while i < n:
    line = lines[i]
    s = line.strip()
    if s == '' or s == '---':
        i += 1; continue
    if s.startswith('# '):
        p = doc.add_heading(level=0); p.add_run(s[2:])
        set_heading_font(p, 18)
        i += 1; continue
    if s.startswith('## '):
        p = doc.add_heading(level=1); p.add_run(s[3:])
        set_heading_font(p, 15)
        i += 1; continue
    if s.startswith('### '):
        p = doc.add_heading(level=2); p.add_run(s[4:])
        set_heading_font(p, 13)
        i += 1; continue
    if s.startswith('#### '):
        p = doc.add_heading(level=3); p.add_run(s[5:])
        set_heading_font(p, 11.5)
        i += 1; continue
    if s.startswith('> '):
        p = doc.add_paragraph(); add_runs(p, s[2:])
        p.runs and setattr(p.runs[0].font, 'italic', True)
        for r in p.runs: r.font.italic = True
        i += 1; continue
    if s.startswith('|'):
        # 收集表格
        tbl_lines = []
        while i < n and lines[i].strip().startswith('|'):
            tbl_lines.append(lines[i].strip()); i += 1
        rows = [re.split(r'\|', ln)[1:-1] for ln in tbl_lines]
        # 去掉分隔行（---）
        data = [r for r in rows if not all(set(c.strip()) <= set('-: ') for c in r)]
        if not data:
            continue
        table = doc.add_table(rows=len(data), cols=len(data[0]))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, rowdata in enumerate(data):
            for ci, cell in enumerate(rowdata):
                c = table.cell(ri, ci)
                c.text = ''
                run = c.paragraphs[0].add_run(cell.strip())
                run.font.size = Pt(9.5)
                if ri == 0:
                    run.bold = True
                    run.font.name = '黑体'; run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        continue
    if re.match(r'^- ', s):
        p = doc.add_paragraph(style='List Bullet'); add_runs(p, s[2:])
        i += 1; continue
    if re.match(r'^\d+\. ', s):
        p = doc.add_paragraph(style='List Number'); add_runs(p, s)
        i += 1; continue
    # 普通段落
    p = doc.add_paragraph(); add_runs(p, s)
    i += 1

# 页脚页码
section = doc.sections[0]
footer = section.footer
fp = footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, 'PAGE')

doc.save(OUT)
print("OK ->", OUT)
